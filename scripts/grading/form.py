"""
Excel coding form for the human reference (LIBRO_CRITERIOS v0.1, section 6).

The coder sees the item number, Task ID, occupation title and task statement
only; K and E are chosen from drop-down lists, each with a one-line reason.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

LEVELS = ["alta", "media", "baja"]

INSTRUCTIONS = [
    ("Codificación humana de referencia: ITEA, codificabilidad (K) y exposición a la IA generativa (E)", "title"),
    ("Versión del libro de criterios: v0.1 (23-10-2026). Codificador único: el autor.", None),
    ("", None),
    ("Cómo trabajar", "h"),
    ("1. Codifica cada tarea con lo que dice el enunciado y el título de la ocupación. No consultes O*NET, ESCO, los grados AXI antiguos ni ningún modelo.", None),
    ("2. Elige K y E en las listas desplegables y escribe en una línea el motivo de cada uno.", None),
    ("3. Codifica K y E de forma independiente: una tarea puede tener K alta y E baja, o al revés.", None),
    ("4. Si dudas entre dos niveles, elige el más probable y marca «Dudosa = sí». No dejes celdas vacías.", None),
    ("5. Trabaja en sesiones de 40–50 tareas, guarda al terminar cada sesión y no vuelvas a tareas ya codificadas.", None),
    ("6. Los enunciados se dejan en inglés, tal como los publica O*NET, para no introducir sesgos de traducción.", None),
    ("", None),
    ("K: codificabilidad del conocimiento", "h"),
    ("Pregunta: ¿el conocimiento que usa la tarea puede expresarse en reglas, procedimientos o representaciones explícitas, o aprenderse de registros documentados, de modo que un sistema de la organización pueda incorporarlo?", None),
    ("ALTA: la tarea trabaja sobre información simbólica (datos, textos, registros, normas) y su criterio puede escribirse como reglas o deducirse de casos documentados. Ejemplo inventado: «Assign billing codes to insurance claims according to coding standards».", None),
    ("MEDIA: combina información documentable con juicio contextual, creatividad o trato con personas; una parte del criterio queda en los registros y otra no. Ejemplo inventado: «Advise managers on restructuring options for underperforming units».", None),
    ("BAJA: depende de destreza física, percepción sensorial en el momento, presencia o una relación personal que se adapta en tiempo real; el criterio no queda en registros. Ejemplo inventado: «Calm and reposition agitated patients during transport».", None),
    ("", None),
    ("E: exposición a la IA generativa", "h"),
    ("Pregunta: ¿un sistema de IA generativa actual (como ChatGPT o Claude, con texto, código o imágenes) puede producir el resultado principal de la tarea, o reducir al menos a la mitad el tiempo que le dedica un trabajador cualificado?", None),
    ("ALTA: sí, directamente; por ejemplo, redactar, resumir, traducir, programar, buscar y sintetizar información o analizar datos descritos en texto.", None),
    ("MEDIA: solo con herramientas o integraciones adicionales (acceso a sistemas, datos o software específico), o solo en una parte relevante de la tarea.", None),
    ("BAJA: no, o de forma marginal; la tarea es física, presencial o sensorial, o su resultado no es un contenido digital.", None),
    ("", None),
    ("Fuentes: la definición de K sigue a Cowan, David y Foray (2000) y a Autor, Levy y Murnane (2003); la de E adapta la rúbrica de Eloundou et al. (2024). Detalle en docs/grading/LIBRO_CRITERIOS_v0.1_ES.md.", None),
]


def write_form(sample, path: str) -> None:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    ins.column_dimensions["A"].width = 150
    for i, (text, style) in enumerate(INSTRUCTIONS, start=1):
        c = ins.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            c.font = Font(bold=True, size=14)
        elif style == "h":
            c.font = Font(bold=True, size=12)

    ws = wb.create_sheet("Codificacion")
    headers = ["N", "Task ID", "Ocupación (O*NET)", "Tarea (O*NET 30.2)",
               "K", "Motivo K", "E", "Motivo E", "Dudosa", "Comentario"]
    widths = [6, 9, 30, 70, 9, 40, 9, 40, 9, 30]
    fill = PatternFill("solid", fgColor="DDE7F5")
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = Font(bold=True), fill
        ws.column_dimensions[c.column_letter].width = w

    rows = sample.sort_values("item")[["item", "Task ID", "Title", "Task"]]
    for i, values in enumerate(rows.itertuples(index=False, name=None), start=2):
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.protection = Protection(locked=True)
        for j in range(5, 11):
            ws.cell(row=i, column=j).protection = Protection(locked=False)

    last = len(rows) + 1
    lv = DataValidation(type="list", formula1='"' + ",".join(LEVELS) + '"', allow_blank=True)
    yn = DataValidation(type="list", formula1='"sí,no"', allow_blank=True)
    for dv in (lv, yn):
        ws.add_data_validation(dv)
    lv.add(f"E2:E{last}")
    lv.add(f"G2:G{last}")
    yn.add(f"I2:I{last}")

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:J{last}"
    # protect the task columns against accidental edits (no password)
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.autoFilter = False
    wb.save(path)
